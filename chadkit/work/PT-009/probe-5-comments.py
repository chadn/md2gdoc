# /// script
# requires-python = ">=3.11"
# dependencies = ["google-api-python-client", "google-auth-oauthlib", "google-auth-httplib2", "python-docx>=1.2"]
# ///
"""Anchored comments and manual formatting under targeted edits and under a whole replace.

The Drive API cannot anchor a comment in a Doc, but a .docx with anchored comments converts to a Doc that has
them, and the .docx export carries each comment's range, so anchoring can be measured without the UI. Five
paragraphs, four with a comment on their first word, one with a red word; then the edits a sync would make."""
import io
import json
import re
import docx
from docx.shared import RGBColor
import h

def anchors(doc_id, label):
    document, comments = h.export_docx(doc_id)
    ranges = h.comment_anchors(document)
    texts = h.comment_texts(comments)
    print(f"--- docx export {label}: {len(texts)} comments, {len(ranges)} anchored")
    for cid, body in sorted(texts.items()):
        print(f"  comment {cid} {body!r} anchored on {ranges.get(cid)!r}")
    print(f"--- drive comments.list {label}")
    for c in h.drive_comments(doc_id):
        print("  ", json.dumps({k: c.get(k) for k in ("content", "anchor", "quotedFileContent", "resolved", "deleted")},
                               ensure_ascii=False))

def red_runs(doc_id):
    doc = h.get(doc_id)
    out = []
    for s, e, t, p in h.paragraphs(doc):
        for r in p["elements"]:
            tr = r.get("textRun")
            if tr and tr.get("textStyle", {}).get("foregroundColor", {}).get("color", {}).get("rgbColor", {}).get("red", 0) > 0.9:
                out.append(tr["content"])
    return out

dx = docx.Document()
spec = [("Alpha", " paragraph, edited within: the word CHANGEME is replaced."),
        ("Beta", " paragraph, deleted entirely."),
        ("Gamma", " paragraph, untouched; a paragraph is inserted after it."),
        ("Epsilon", " paragraph, whole text replaced by delete plus insert.")]
for first, rest in spec:
    p = dx.add_paragraph()
    r1 = p.add_run(first)
    p.add_run(rest)
    dx.add_comment(runs=[r1], text=f"comment on {first}", author="probe", initials="pr")
p = dx.add_paragraph()
p.add_run("Delta paragraph with a ")
r = p.add_run("red word")
r.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
p.add_run(" set by hand, untouched.")
buf = io.BytesIO()
dx.save(buf)
f = h.create_from(buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                  "PT-009 probe 5 comments")
doc_id = f["id"]
print("doc", doc_id, f["webViewLink"])
open("comments-doc-id.txt", "w").write(doc_id)
doc = h.get(doc_id)
h.show(doc, "as converted")
print("red runs:", red_runs(doc_id))
anchors(doc_id, "as converted")

# targeted edits, bottom-up so earlier indexes hold within one batch
doc = h.get(doc_id)
a = h.find(doc, "Alpha")
word = a[0] + a[2].index("CHANGEME")
b = h.find(doc, "Beta")
g = h.find(doc, "Gamma")
e = h.find(doc, "Epsilon")
reqs = [
    {"deleteContentRange": {"range": {"startIndex": e[0], "endIndex": e[1] - 1}}},
    {"insertText": {"location": {"index": e[0]}, "text": "Epsilon rewritten by the API."}},
    {"insertText": {"location": {"index": g[1] - 1}, "text": "\nNew paragraph after Gamma."}},
    {"deleteContentRange": {"range": {"startIndex": b[0], "endIndex": b[1]}}},
    {"deleteContentRange": {"range": {"startIndex": word, "endIndex": word + len("CHANGEME")}}},
    {"insertText": {"location": {"index": word}, "text": "CHANGED"}},
]
h.batch(doc_id, reqs)
h.show(h.get(doc_id), "after targeted edits")
print("red runs:", red_runs(doc_id))
anchors(doc_id, "after targeted edits")
print("--- markdown export after targeted edits")
md = h.export_md(doc_id)
print(md)

# the script's whole replace with the same text
h.d.update(doc_id, h.md2gdoc.normalize(md))
h.show(h.get(doc_id), "after whole replace")
print("red runs:", red_runs(doc_id))
anchors(doc_id, "after whole replace")
