# /// script
# requires-python = ">=3.11"
# dependencies = ["google-api-python-client", "google-auth-oauthlib", "google-auth-httplib2", "python-docx>=1.2"]
# ///
"""Can a push travel through a .docx body and keep what the markdown body drops? A Doc from markdown with a
ticked box, a comment anchored by the Docs API and a red word; its .docx export sent back unchanged through
files.update, then edited with python-docx (a word changed, a box ticked, a paragraph added) and sent again. After
each step: the markdown export, the comment anchors, the red word."""
import io
import re
import docx
from googleapiclient.http import MediaIoBaseUpload
import h

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
SRC = """# Carrier

- [ ] unticked
- [x] ticked
- [ ] tick me via docx

A paragraph with **bold** and the word CHANGEME inside, and a red word later.

Commented paragraph stays as it is.
"""
f = h.d.create("PT-009 probe 11 docx carrier", h.md2gdoc.push_form(SRC), None)
doc_id = f["doc_id"]
print("doc", doc_id)
open("carrier-doc-id.txt", "w").write(doc_id)
doc = h.get(doc_id)
c = h.find(doc, "Commented")
p = h.find(doc, "A paragraph")
r = p[0] + p[2].index("red word")
h.batch(doc_id, [
    {"insertComment": {"content": "anchored on Commented", "range": {"startIndex": c[0], "endIndex": c[0] + 9}}},
    {"updateTextStyle": {"range": {"startIndex": r, "endIndex": r + 8},
                         "textStyle": {"foregroundColor": {"color": {"rgbColor": {"red": 1}}}}, "fields": "foregroundColor"}},
])

def state(label):
    print(f"\n--- {label}")
    print(h.export_md(doc_id))
    document, comments = h.export_docx(doc_id)
    print("docx comment anchors:", h.comment_anchors(document))
    doc = h.get(doc_id)
    red = [tr["content"] for *_, p in h.paragraphs(doc) for pe in p["elements"]
           if (tr := pe.get("textRun")) and tr.get("textStyle", {}).get("foregroundColor", {}).get("color", {}).get("rgbColor", {}).get("red") == 1]
    print("red runs:", red, "| Drive version:", h.d.version(doc_id))
    return document

document = state("as made")
raw = h.drive.files().export(fileId=doc_id, mimeType=DOCX).execute()
open("carrier-export.docx", "wb").write(raw)
# how the docx represents the ticked and the unticked item
for needle in ("unticked", "ticked<", "tick me"):
    m = re.search(r"<w:p\b(?:(?!<w:p\b).)*?%s.*?</w:p>" % re.escape(needle), document, re.S)
    print(f"\ndocx paragraph holding {needle!r}:")
    print(re.sub(r"<w:rPr>.*?</w:rPr>", "<w:rPr/>", m.group(0))[:700] if m else "not found")
numbering = h.zipfile.ZipFile(io.BytesIO(raw)).read("word/numbering.xml").decode()
print("\nnumbering.xml, the checkbox glyphs:")
for m in re.finditer(r'<w:lvlText w:val="([^"]*)"/>', numbering):
    v = m.group(1)
    print("  lvlText", repr(v), [hex(ord(ch)) for ch in v][:3])
print("w14:checkbox present:", "checkbox" in document)

def update(data, label):
    h.drive.files().update(fileId=doc_id, media_body=MediaIoBaseUpload(io.BytesIO(data), mimetype=DOCX),
                           fields="version", supportsAllDrives=True).execute()
    return state(label)

update(raw, "after files.update with the unchanged .docx")

dx = docx.Document(io.BytesIO(raw))
for para in dx.paragraphs:
    for run in para.runs:
        if "CHANGEME" in run.text:
            run.text = run.text.replace("CHANGEME", "CHANGED")
ticked = [p for p in dx.paragraphs if p.text.strip() == "ticked"][0]
target = [p for p in dx.paragraphs if p.text.strip() == "tick me via docx"][0]
# copy the ticked item's paragraph properties (numbering, run glyph) onto the target
t_ppr = ticked._p.pPr
target._p.remove(target._p.pPr) if target._p.pPr is not None else None
target._p.insert(0, h.copy.deepcopy(t_ppr))
last = dx.paragraphs[-1]
new = last.insert_paragraph_before("Added through the docx.")
buf = io.BytesIO()
dx.save(buf)
update(buf.getvalue(), "after the edited .docx (CHANGEME->CHANGED, third box given the ticked item's numbering, a paragraph added)")
